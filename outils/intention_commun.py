# -*- coding: utf-8 -*-
"""Les fonctions communes des fiches d'intention pédagogique (une par scénario Simulog).

Règle (fiche projet `prepalog-fiche-intention-pedagogique`, décision 17 du 03/10/2026) : chaque scénario a
UNE fiche d'intention, pour l'enseignant (Tristan, son équipe, un remplaçant) ; Word + PDF générés par Cowork,
rangés dans `contenus/intentions/`, ouverts par le bouton « Fiche d'intention » (enseignant seul) une fois la
ligne `intention` déclarée dans `ENTREPRISES` (`activites/index.js`) — déclarer, c'est valider.

Le fichier est PUBLIC (comme tout le dépôt) : aucun secret, jamais la valeur d'un code d'accès.

Mise en page : Calibri, noir et gris (seuls les logos portent la couleur), comme les trames. Le PDF est
produit par LibreOffice (`finir(..., pdf=True)`).

Usage :
    import intention_commun as I
    I.nouveau()
    I.couverture([logo1, logo2], titre, sous_titre, encadre)
    I.partie('A. Le scénario') ; I.h2('…') ; I.p('…') ; I.puces([...]) ; I.tableau([...], [[...]])
    I.seance('ENT-5.1', 'titre', etat) …
    I.finir('smoby')
"""
import os, subprocess
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ENCRE = RGBColor(0x1a, 0x1a, 0x1a)
TITRE = RGBColor(0x11, 0x11, 0x11)
GRIS = RGBColor(0x59, 0x59, 0x59)

ICI = os.path.dirname(os.path.abspath(__file__))
RACINE = os.path.join(ICI, '..')
LOGOS = os.path.join(RACINE, 'contenus', 'trames', 'logos')
SORTIE = os.path.join(RACINE, 'contenus', 'intentions')

d = None


def nouveau():
    global d
    d = Document()
    st = d.styles['Normal']
    st.font.name = 'Calibri'; st.font.size = Pt(10.5); st.font.color.rgb = ENCRE
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.08
    for s in d.sections:
        s.top_margin = Cm(1.4); s.bottom_margin = Cm(1.4)
        s.left_margin = s.right_margin = Cm(1.8)
    _pied()
    return d


def _pied():
    """Pied de page : nom de la fiche à gauche, numéro de page."""
    s = d.sections[0]
    par = s.footer.paragraphs[0]
    r = par.add_run('Prepalog · Simulog — fiche d\'intention pédagogique (enseignant)    ·    page ')
    r.font.size = Pt(8); r.font.color.rgb = GRIS
    r = par.add_run(); r.font.size = Pt(8); r.font.color.rgb = GRIS
    for typ, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if typ:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), typ); r._r.append(e)
        else:
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r._r.append(e)


def ombre(cell, hexa):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexa)
    tcPr.append(sh)


def _riche(par, texte, taille=None, couleur=None):
    """**gras** dans le texte."""
    morceaux = texte.split('**')
    for i, m in enumerate(morceaux):
        if not m:
            continue
        r = par.add_run(m); r.bold = (i % 2 == 1)
        if taille: r.font.size = Pt(taille)
        if couleur: r.font.color.rgb = couleur
    return par


def p(texte='', taille=None, couleur=None, apres=4, avant=0, gras=False, italique=False):
    par = d.add_paragraph(); par.paragraph_format.space_after = Pt(apres); par.paragraph_format.space_before = Pt(avant)
    if gras or italique:
        r = par.add_run(texte); r.bold = gras; r.italic = italique
        if taille: r.font.size = Pt(taille)
        if couleur: r.font.color.rgb = couleur
    else:
        _riche(par, texte, taille, couleur)
    return par


def saut():
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def partie(titre):
    par = d.add_paragraph(); par.paragraph_format.space_before = Pt(4); par.paragraph_format.space_after = Pt(8)
    par.paragraph_format.keep_with_next = True
    r = par.add_run(titre.upper()); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = TITRE
    _filet(par)


def _filet(par):
    pPr = par._p.get_or_add_pPr(); bd = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '8'); b.set(qn('w:space'), '2'); b.set(qn('w:color'), '595959')
    bd.append(b); pPr.append(bd)


def h2(titre):
    par = d.add_paragraph(); par.paragraph_format.space_before = Pt(8); par.paragraph_format.space_after = Pt(3)
    par.paragraph_format.keep_with_next = True
    r = par.add_run(titre); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = TITRE


def h3(titre):
    par = d.add_paragraph(); par.paragraph_format.space_before = Pt(5); par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.keep_with_next = True
    r = par.add_run(titre); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = GRIS


def puces(items, numeros=False):
    for i, t in enumerate(items, 1):
        par = d.add_paragraph(); par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.left_indent = Cm(0.6); par.paragraph_format.first_line_indent = Cm(-0.45)
        r = par.add_run(f'{i}. ' if numeros else '•  '); r.bold = numeros
        _riche(par, t)


def tableau(entetes, lignes, largeurs=None, taille=9.5):
    t = d.add_table(rows=1, cols=len(entetes)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, e in enumerate(entetes):
        c = t.rows[0].cells[i]; ombre(c, 'E7E7E7')
        par = c.paragraphs[0]; par.paragraph_format.space_after = Pt(0)
        r = par.add_run(e); r.bold = True; r.font.size = Pt(taille)
    _entete_repete(t.rows[0])
    for ligne in lignes:
        row = t.add_row(); row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i, v in enumerate(ligne):
            par = row.cells[i].paragraphs[0]; par.paragraph_format.space_after = Pt(0)
            _riche(par, str(v), taille)
    if largeurs:
        largeurs_fixes(t, largeurs)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def largeurs_fixes(t, largeurs):
    """Largeurs de colonnes respectées par Word ET LibreOffice (grille + disposition fixe)."""
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        if i < len(largeurs):
            gc.set(qn('w:w'), str(int(largeurs[i].twips)))
    for row in t.rows:
        for i, w in enumerate(largeurs):
            row.cells[i].width = w


def _entete_repete(row):
    trPr = row._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)


def encadre(titre, lignes, fond='F2F2F2'):
    t = d.add_table(rows=1, cols=1); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.LEFT
    row = t.rows[0]; row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    c = row.cells[0]; ombre(c, fond)
    par = c.paragraphs[0]; par.paragraph_format.space_after = Pt(3)
    r = par.add_run(titre); r.bold = True; r.font.size = Pt(10.5)
    for l in lignes:
        o = c.add_paragraph(); o.paragraph_format.space_after = Pt(2)
        o.paragraph_format.left_indent = Cm(0.45); o.paragraph_format.first_line_indent = Cm(-0.35)
        o.add_run('•  ')
        _riche(o, l, 10)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def couverture(logos, titre, sous_titre, lignes_etat):
    """Page 1 : logos côte à côte, titre, sous-titre, encadré « état de la fiche »."""
    par = d.add_paragraph(); par.paragraph_format.space_after = Pt(10)
    for chemin, hauteur in logos:
        par.add_run().add_picture(chemin, height=hauteur)
        par.add_run('      ')
    par = d.add_paragraph(); par.paragraph_format.space_after = Pt(2)
    r = par.add_run('FICHE D\'INTENTION PÉDAGOGIQUE — à l\'usage de l\'enseignant'); r.font.size = Pt(9); r.font.color.rgb = GRIS
    r.bold = True
    par = d.add_paragraph(); par.paragraph_format.space_after = Pt(2)
    r = par.add_run(titre); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = TITRE
    par = d.add_paragraph(); par.paragraph_format.space_after = Pt(10)
    r = par.add_run(sous_titre); r.font.size = Pt(11.5); r.font.color.rgb = GRIS
    encadre('État de cette fiche', lignes_etat)


def seance(code, titre, etat, nouvelle_page=True):
    """Ouvre une séance (sur une nouvelle page par défaut) : code + titre, puis la mention d'état."""
    der = d.element.body[-2] if d.element.body[-1].tag == qn('w:sectPr') else d.element.body[-1]
    if nouvelle_page and der.tag == qn('w:p') and not ''.join(der.itertext()).strip():
        der.getparent().remove(der)   # le paragraphe vide après un tableau, qui ferait une page blanche
    par = d.add_paragraph(); par.paragraph_format.space_after = Pt(1); par.paragraph_format.keep_with_next = True
    # Saut porté par le titre lui-même (pas de paragraphe vide qui pourrait déborder et laisser une page blanche).
    par.paragraph_format.page_break_before = nouvelle_page
    r = par.add_run(code + '  '); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = GRIS
    r = par.add_run(titre); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = TITRE
    _filet(par)
    par = d.add_paragraph(); par.paragraph_format.space_after = Pt(6)
    r = par.add_run(etat); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GRIS


def fiche_identite(lignes):
    """Petit tableau à deux colonnes (libellé gris, valeur)."""
    t = d.add_table(rows=0, cols=2); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for lib, val in lignes:
        row = t.add_row(); row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        c0, c1 = row.cells; ombre(c0, 'F2F2F2')
        a = c0.paragraphs[0]; a.paragraph_format.space_after = Pt(0)
        rr = a.add_run(lib); rr.bold = True; rr.font.size = Pt(9.5)
        b = c1.paragraphs[0]; b.paragraph_format.space_after = Pt(0); _riche(b, val, 9.5)
    largeurs_fixes(t, [Cm(4.2), Cm(13.2)])
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def finir(nom, pdf=True):
    """Écrit contenus/intentions/<nom>-intention-pedagogique.docx (et .pdf avec LibreOffice)."""
    os.makedirs(SORTIE, exist_ok=True)
    for t in d.element.body.iter(qn('w:t')):
        if t.text and ('« ' in t.text or ' »' in t.text):
            t.text = t.text.replace('« ', '« ').replace(' »', ' »')
    base = os.path.join(SORTIE, f'{nom}-intention-pedagogique')
    d.save(base + '.docx')
    print('fiche :', base + '.docx')
    if pdf:
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', SORTIE, base + '.docx'],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        print('pdf   :', base + '.pdf')
