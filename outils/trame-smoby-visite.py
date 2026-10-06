# -*- coding: utf-8 -*-
"""Trame élève d'ENT-5.3 Smoby, « la visite de la plateforme » (Cowork, 06/10/2026 au soir) — BROUILLON.

Format LONG (6 étapes + feuille de cours à détacher), comme ENT-5.1 et 5.2 (décision de Tristan du 06/10/2026).
Écrite AVANT la validation à l'écran, à la demande de Tristan (« toute la suite 5.3 → 5.8 »), d'après l'écran TEL
QU'IL SERA APRÈS LE LOT A de `docs/briefs/SMOBY-retours-5.3.md` (voix au tu, point 7 « Travée » et point 8
« Croisillons » déplacés, échelle et lisse sans couleur). Les passages marqués « À REVOIR APRÈS LE LOT A » sont à
relire quand Claude Code l'aura livré.

Vérifié : la plateforme de stockage de Smoby à Moirans-en-Montagne (hebdo39.net) ; la règle de circulation
piétons / engins (séparer les flux, passages piétons matérialisés : INRS ED 6350) ; la plaque de charge
des racks (INRS ED 771, NF EN 15635).
Construit (comme dans la séance) : Bruno, le parcours, le plan, le stock, l'adresse A1-T03-N2-E1.
Données relues dans `contenus/smoby-ent53.js` (points, parcours, travée, adresse) et `contenus/smoby-entrepot.js`
(plan et stock, pour les adresses de l'étape 5 : le script vérifie qu'elles existent et lit ce qu'il y a dedans).

Photos : les mêmes qu'à l'écran (libres, Pexels ; dessin de Cowork), imprimées en NIVEAUX DE GRIS (règle des trames :
seuls les logos en couleur), numérotées ou lettrées par ce script en mémoire (aucune image nouvelle dans le dépôt).

Lancer : python3 outils/trame-smoby-visite.py
puis    soffice --headless --convert-to pdf --outdir contenus/trames contenus/trames/ENT-5.3-smoby-visite-trame-eleve.docx
"""
import io, os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import trame_commun as T
import corriges_data
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont

LOGO = os.path.join(T.RACINE, 'contenus', 'trames', 'logos', 'smoby.png')
VISITE = os.path.join(T.RACINE, 'contenus', 'smoby', 'visite')
POLICE = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'


# ==================================================================== photos (gris, repères dessinés)
def photo(fichier, largeur_cm, reperes=(), echelle=1.0, mention='', rayon=30, taille=34):
    """La photo en niveaux de gris, avec des pastilles blanches (`reperes` : (texte, x, y) dans le repère de
    l'écran, multiplié par `echelle` pour tomber dans le fichier), centrée, puis la mention en petit."""
    im = Image.open(os.path.join(VISITE, fichier)).convert('L').convert('RGB')
    dr = ImageDraw.Draw(im)
    f = ImageFont.truetype(POLICE, taille)
    for txt, x, y in reperes:
        x, y = x * echelle, y * echelle
        dr.ellipse([x - rayon, y - rayon, x + rayon, y + rayon], fill='white', outline='black', width=4)
        dr.text((x, y), txt, fill='black', font=f, anchor='mm')
    im.thumbnail((1100, 1100))   # assez pour l'impression à cette taille ; la trame reste légère
    buf = io.BytesIO(); im.save(buf, 'JPEG', quality=78); buf.seek(0)
    T.colle_au_suivant()
    par = T.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(0); par.paragraph_format.keep_with_next = True
    if T._consomme(): par.paragraph_format.page_break_before = True
    par.add_run().add_picture(buf, width=Cm(largeur_cm))
    par = T.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(6)
    r = par.add_run(mention); r.font.size = Pt(8); r.font.color.rgb = T.GRIS


# ==================================================================== données (relues dans contenus/smoby-ent53.js)
# Vue du ciel (repère 1600 × 1066 = le fichier) : les 6 points de l'écran, lettrés ici DANS UN AUTRE ORDRE que les
# numéros de l'écran (l'élève ne recopie pas : il reconnaît).
CIEL = {1: ('L’entrepôt', 1080, 120), 2: ('Les quais', 470, 420), 3: ('Parking poids lourds', 1040, 575),
        4: ('Aire de manœuvre', 330, 760), 5: ('Passage piétons', 440, 578), 6: ('Parking des salariés', 667, 1000)}
LETTRES = {'A': 4, 'B': 1, 'C': 6, 'D': 2, 'E': 5, 'F': 3}
QUI_CIRCULE = {1: 'Les chariots et les salariés à pied (à l’intérieur).', 2: 'Les camions, à reculons contre la porte.',
               3: 'Les camions qui attendent leur tour.', 4: 'Les camions qui manœuvrent : pas de piétons.',
               5: 'Les piétons : seul chemin pour traverser la cour.', 6: 'Les voitures des salariés.'}
# Le parcours (dans l'ordre) et ce qu'en dit Bruno.
PARCOURS = ['Le quai', 'La zone de réception', 'L’allée principale', 'Les allées de stockage', 'La zone litiges',
            'Le bureau du chef de quai']
DESORDRE = [4, 2, 0, 5, 3, 1]
CE_QU_ON_Y_FAIT = ['Les camions arrivent à reculons ; on décharge un camion calé.',
                   'Les palettes déchargées attendent ; on les contrôle avec le bon de livraison.',
                   'Les chariots roulent dans les deux sens ; à pied, on reste sur le côté.',
                   'Les racks à palettes ; chaque emplacement a une adresse sur la lisse.',
                   'Les palettes abîmées ou en attente du fournisseur, en L1 ou L2 ; pas en stock.',
                   'Un problème, un document à signer : on va voir Bruno.']
# Les mots du rack (photo `allee`, repère 1400 × 788, fichier 1600 × 900). À REVOIR APRÈS LE LOT A : points 7 et 8
# à leur place nouvelle (A2 du brief SMOBY-retours-5.3), définitions sans couleur (A3).
MOTS = [(1, 215, 430, 'Échelle', 'Le montant vertical, percé de trous, qui porte les lisses.'),
        (2, 430, 247, 'Lisse', 'La barre horizontale sur laquelle on pose les palettes.'),
        (3, 155, 182, 'Étiquette d’adresse', 'L’adresse de l’emplacement, collée sur la lisse.'),
        (4, 385, 470, 'Palette filmée', 'Des cartons tenus par un film plastique étirable.'),
        (5, 740, 620, 'Allée', 'Le couloir entre deux racks, où roulent les chariots.'),
        (6, 1300, 290, 'Niveau', 'Un étage de lisses. Le sol est le niveau 1.'),
        (7, 330, 330, 'Travée', 'L’espace entre deux échelles, sur toute la hauteur.'),
        (8, 1100, 420, 'Croisillons', 'Les barres en diagonale qui rigidifient l’échelle.')]

# Le plan et le stock (relus dans contenus/smoby-entrepot.js) : de quoi juger les adresses de l'étape 5.
def _lire_js(chemin):
    with open(chemin, encoding='utf-8') as f:
        return f.read()
_SRC = _lire_js(os.path.join(T.RACINE, 'contenus', 'smoby-entrepot.js'))
STOCK = {a: (p, int(kg)) for a, p, kg in re.findall(r"'([AB][12]-T0\d-N\d-E\d)': S\('([A-Z]{3})', (\d+)\)", _SRC)}
NOMS = dict(re.findall(r"^\s+([A-Z]{3}): \{ nom: '([^']+)'", _SRC, re.M))
assert len(STOCK) == 99, len(STOCK)                         # le stock figé de la maquette v2
assert STOCK['A1-T03-N2-E1'] == ('MAI', 420)               # la cible de la séance
ADRESSE = 'A1-T03-N2-E1'


def adresse(cote, t, n, e):
    return f'{cote}-T{t:02d}-N{n}-E{e}'


def decompose(a):
    m = re.fullmatch(r'([AB])([12])-T(\d\d)-N(\d)-E(\d)', a)
    al, co, t, n, e = m.groups()
    return [f'allée {al}, côté {co}', str(int(t)), n + (' (le sol)' if n == '1' else ''), e]


# Étape 5 « À ton tour » : deux adresses à décomposer, deux à écrire. Toutes existent dans le stock (vérifié).
A_LIRE = ['A2-T01-N1-E3', 'B1-T02-N3-E2']
A_ECRIRE = [('B2', 4, 1, 3), ('A1', 1, 3, 2)]
for a in A_LIRE + [adresse(*x) for x in A_ECRIRE]:
    assert a in STOCK, a


# ==================================================================== la trame
T.nouveau()
T.entete(LOGO, 'ENT-5.3 — Carnet de suivi : la visite de la plateforme', [
    ('Ce document est ta trame de travail :', "tu peux le suivre seul, étape par étape. Aujourd'hui, tu changes de "
     "poste : tu es Yanis, le cariste que tu as recruté. C'est ton premier jour. Avant de toucher un chariot, Bruno, "
     "le chef de quai, te fait visiter la plateforme."),
    ('Ce que ton enseignant voit dans son suivi :', "dix-sept points : les 3 questions de la vue du ciel, les 6 photos "
     "replacées sur le plan, les 4 éléments du quiz, la travée (2) et l'adresse (2). Tes réponses écrites ici servent "
     "à réfléchir et à retenir."),
    ('Ce qui est vrai, ce qui est inventé :', "la plateforme de Smoby à Moirans-en-Montagne (Jura) et les règles de "
     "sécurité sont réelles. Bruno, le plan, le parcours et les adresses sont inventés pour l'exercice. Les photos "
     "viennent d'autres entrepôts.")],
    [('Le message de Bruno et la vue du ciel', 'Prepalog : Messagerie et Visite'),
     ('Le parcours d’une palette', 'Prepalog : Visite (Le parcours, Où est-ce ?)'),
     ('Les mots du rack', 'Prepalog : Visite (Les mots du rack, Quiz)'),
     ('La travée', 'Prepalog : Visite (La travée)'),
     ('L’adresse d’un emplacement', 'Prepalog : Visite (L’adresse) et ce carnet'),
     ('Fin de la visite : la sécurité', 'Prepalog : Visite (Fin)')], nom_logo='SMOBY')

# ==================================================================== étape 1
T.etape(1, 'Le message de Bruno et la vue du ciel')
T.consignes(["Ouvre la Messagerie et lis le message de Bruno.",
             "Ouvre le menu « Visite de la plateforme ». Lis l'accueil, puis clique « Suivant ».",
             "Vue du ciel : ouvre les 6 points, puis réponds aux 3 questions en cliquant sur la photo."])
T.p("Sur cette photo, les lettres sont à la place des 6 points de l'écran, mais pas dans le même ordre. Écris le nom "
    "de chaque endroit, puis qui y circule. Trace aussi au crayon le chemin d'un salarié qui va à pied de sa "
    "voiture jusqu'à l'entrepôt.", apres=4)
photo('ciel-pexels-2804929.jpg', 8.6, [(l, CIEL[n][1], CIEL[n][2]) for l, n in LETTRES.items()], rayon=34, taille=38,
      mention='Photo d’un autre entrepôt : Marcin Jozwiak, Pexels n° 2804929.')
T.tableau(['Lettre', 'Nom de l’endroit', 'Qui y circule ?'], 0, [Cm(1.6), Cm(5.6), Cm(9.8)], hauteur=Cm(0.55),
          remplis=[[l] for l in LETTRES])
T.faits(['Qui est Bruno ?', 'Quel jour commence Yanis, et à quelle heure ?',
         'Par où ton chemin traverse-t-il la cour des camions ?'], hauteur=Cm(0.65))
T.reflechir(["Sur cette plateforme, le parking des salariés est loin de celui des camions. Pourquoi ?"])

# ==================================================================== étape 2
T.etape(2, 'Le parcours d’une palette')
T.consignes(["Clique les 6 étapes du parcours dans l'ordre, sur le plan. Lis chaque fois ce que dit Bruno.",
             "Remplis le tableau : le rang de chaque endroit dans le parcours, et ce qu'on y fait.",
             "Étape « Où est-ce ? » : replace chaque photo sur le plan."])
T.tableau(['Endroit (dans le désordre)', 'Rang (1 à 6)', 'Ce qu’on y fait'], 0, [Cm(4.6), Cm(2.2), Cm(10.2)],
          hauteur=Cm(0.72), remplis=[[PARCOURS[i]] for i in DESORDRE])
T.faits(['Que fait le niveleur, au quai ?', 'À la réception, avec quel document contrôle-t-on les palettes ?'],
        hauteur=Cm(0.7))
T.p("Ce dessin montre ce qu'on doit voir dans une zone litiges.", apres=2)
photo('visite-litiges-dessin.jpg', 5.8, mention='Dessin : ce qu’on doit voir dans une zone litiges.')
T.faits(['Cite deux signes qui montrent qu’une palette de cette zone ne doit pas partir.',
         'Pourquoi la palette L2 est-elle bloquée ?'], hauteur=Cm(0.7))
T.qcm([("Bruno le dit au quai : on ne décharge jamais un camion…",
        ['qui n’est pas calé', 'qui est en retard', 'qui vient d’Arinthod'], 0)])
T.reflechir(["Pourquoi une palette abîmée ne va-t-elle pas directement en stock ?"])

# ==================================================================== étape 3
T.etape(3, 'Les mots du rack')
T.consignes(["Ouvre les 8 mots du rack à l'écran et lis-les.",
             "Sans regarder l'écran, écris le mot de chaque numéro dans le tableau. Puis vérifie à l'écran et "
             "corrige d'une autre couleur.",
             "Fais le quiz : retrouve les éléments sur la photo d'un autre entrepôt."])
photo('allee-pexels-5775099.jpg', 9.4, [(str(n), x, y) for n, x, y, _, _ in MOTS], echelle=1600 / 1400,
      mention='Photo d’un autre entrepôt : Handi Boyz LLC, Pexels n° 5775099.')
T.tableau(['N°', 'Mot', 'Ce que c’est, en quelques mots'], 0, [Cm(1.2), Cm(4.4), Cm(11.4)], hauteur=Cm(0.55),
          remplis=[[str(n)] for n, *_ in MOTS])
T.qcm([("Un croisillon est tordu (un chariot l'a heurté). Que fais-tu ?",
        ['rien, il tient', 'je le signale au chef de quai', 'je le redresse'], 1),
       ("Le sol est le niveau…", ['0', '1', '3'], 1)])
T.reflechir(["Pourquoi ne faut-il pas reconnaître une échelle à sa couleur ? Comment la reconnais-tu, alors ?"])

# ==================================================================== étape 4
T.etape(4, 'La travée')
T.consignes(["À l'écran, place les 4 coins de la travée complète, puis clique ses lisses.",
             "Sur cette photo, entoure au crayon la même travée, puis fais une croix sur chacune de ses lisses."])
photo('visite-travee-pexels-29454378.jpg', 8.6,
      mention='Photo d’un autre entrepôt : Pexels n° 29454378, recadrée, marques floutées.')
T.faits(['Combien de lisses compte cette travée ?', 'Entre quoi et quoi va une travée, en largeur ? Et en hauteur ?',
         'Le bas de cette travée est un passage, sans palette. Est-ce quand même une travée ?',
         'Les barres qu’on voit au fond, à travers la travée, en font-elles partie ? Pourquoi ?'], hauteur=Cm(0.85))
T.qcm([("Une travée, c'est…", ['une seule palette', 'l’espace entre deux échelles, sur toute la hauteur du rack',
                               'le couloir où roulent les chariots'], 1)])
T.reflechir(["Bruno te dit « Va en travée 3 » plutôt que « Va à la troisième étagère ». Pourquoi est-ce plus sûr ?"])

# ==================================================================== étape 5
T.etape(5, 'L’adresse d’un emplacement')
T.encadre('Lire une adresse :', "A1 = allée A, côté 1 (une allée a deux côtés : A1 et A2 sont les racks de part et "
          "d'autre de l'allée A). T = travée, N = niveau (le sol est N1), E = emplacement (3 palettes par niveau).")
T.consignes(["Bruno te montre l'étiquette " + ADRESSE + ". Décompose-la ici, au crayon.",
             "À l'écran, choisis le sens de chaque partie et valide : tu n'as qu'un essai. Vérifie avant.",
             "Retrouve l'emplacement : la travée sur le plan, puis l'emplacement dans la travée vue de face."])
T.tableau(['Partie de l’adresse', 'Ce qu’elle veut dire'], 0, [Cm(4.6), Cm(12.4)], hauteur=Cm(0.75),
          remplis=[[x] for x in ADRESSE.split('-')])
T.faits(['Qu’as-tu trouvé à l’emplacement ' + ADRESSE + ' ?'], hauteur=Cm(0.8))
T.p("À ton tour. Ces emplacements sont sur la même plateforme. Complète chaque ligne.", apres=4)
T.tableau(['Adresse', 'Allée et côté', 'Travée', 'Niveau', 'Emplacement'], 0,
          [Cm(4.2), Cm(3.8), Cm(2.6), Cm(3.2), Cm(3.2)], hauteur=Cm(0.85),
          remplis=[[a] for a in A_LIRE] + [['', f'allée {c[0]}, côté {c[1]}', str(t), str(n), str(e)]
                                         for c, t, n, e in A_ECRIRE])
T.qcm([("À l'adresse " + ADRESSE + ", la palette est posée…",
        ['au sol', 'sur les premières lisses, juste au-dessus du sol', 'tout en haut du rack'], 1)])
T.reflechir(["Pourquoi une adresse se lit-elle de la plus grande zone (l'allée) à la plus petite (l'emplacement) ?"])

# ==================================================================== étape 6
T.etape(6, 'Fin de la visite : la sécurité')
T.consignes(["Lis le dernier message de Bruno (étape « Fin »).",
             "Réponds aux questions : elles préparent ta prochaine séance, au quai."])
T.faits(['Que fera Yanis cet après-midi ?', 'Par quoi Bruno dit-il qu’on commence ?'], hauteur=Cm(0.9))
T.encadre('À savoir :', "sur une plateforme, piétons et chariots sont séparés autant que possible : allées marquées "
          "au sol, passages piétons, priorité aux règles du site. Sur chaque rack, une plaque indique la charge "
          "maximale qu'on peut poser : on ne la dépasse jamais.")
T.qcm([("L'aire de manœuvre est le grand espace où les camions reculent vers les quais. Un piéton…",
        ['la traverse s’il se dépêche', 'n’y circule jamais', 'y marche au milieu'], 1),
       ("Où est écrite la charge maximale qu'on peut poser sur une lisse ?",
        ['sur une plaque, au bout du rack', 'sur le bon de livraison', 'nulle part : on regarde si ça plie'], 0),
       ("À pied dans l'allée principale, tu…",
        ['marches au milieu', 'restes sur le côté et traverses au passage piétons', 'cours pour gêner moins longtemps'], 1)])
T.reflechir(["Tu as vu six endroits ce matin. Lequel te paraît le plus dangereux pour un piéton ? Pourquoi ?",
             "Qu'est-ce qui t'aidera le plus cet après-midi, au quai : le plan, les mots du rack ou les adresses ? "
             "Pourquoi ?"])

# ==================================================================== feuille à détacher (cours)
ESSENTIEL = [("Sur une plateforme, piétons et camions ne se croisent pas : à pied, on traverse la cour au {}.",
              'passage piétons'),
             ('Une {} est l’espace entre deux échelles, sur toute la hauteur du rack.', 'travée'),
             ("Une adresse se lit de la plus grande zone à la plus petite : allée et côté, travée, {}, emplacement.",
              'niveau'),
             ('Une palette abîmée ne va pas en stock : elle attend en zone {}.', 'litiges')]
LEXIQUE = [('échelle', 'Montant vertical, percé de {}, qui porte les lisses.', 'trous'),
           ('lisse', 'Barre {} sur laquelle on pose les palettes.', 'horizontale'),
           ('quai', 'Porte où un camion se met à {} pour être chargé ou déchargé.', 'reculons'),
           ('croisillons', 'Barres en {} qui rigidifient l’échelle.', 'diagonale')]
T.feuille_cours('ENT-5.3', 'Se repérer sur une plateforme logistique',
                'Logistique — C1.2 règles de sécurité, C1.5 mettre en stock (initiation)',
                [ph.format(T.TROU) for ph, _ in ESSENTIEL],
                [m for _, m in ESSENTIEL] + [m for _, _, m in LEXIQUE],
                [(mot, df.format(T.TROU)) for mot, df, _ in LEXIQUE])

# ==================================================================== corrigé de la trame
ENT_5_3 = {
 "Qui est Bruno": {"rep": "Le chef de quai de la plateforme Smoby."},
 "Quel jour commence Yanis": {"rep": "Le mercredi 9 décembre, à 8 h."},
 "T: Lettre | Nom de l’endroit": {"lignes": [[l, CIEL[n][0], QUI_CIRCULE[n]] for l, n in LETTRES.items()],
   "note": "Lettres ≠ numéros de l'écran (A = point 4, B = 1, C = 6, D = 2, E = 5, F = 3). Accepter toute formulation équivalente pour « qui y circule »."},
 "Par où ton chemin traverse-t-il": {"rep": "Par le passage piétons (lettre E), jamais par l'aire de manœuvre (A)."},
 "Sur cette plateforme, le parking des salariés": {"pistes": [
   "Pour que voitures et camions ne se croisent pas : un camion qui recule voit mal.",
   "Un piéton qui sort de sa voiture ne doit pas se retrouver au milieu des camions.",
   "Chacun a sa zone : moins de risques d'accident, et la circulation est plus simple."]},
 "T: Endroit (dans le désordre) | Rang": {"lignes": [[PARCOURS[i], str(i + 1), CE_QU_ON_Y_FAIT[i]] for i in DESORDRE],
   "note": "Ordre du parcours : " + ' → '.join(PARCOURS) + ". Le chemin d'une palette, du quai jusqu'au rack (puis les litiges et le bureau)."},
 "Que fait le niveleur": {"rep": "Il fait le pont entre le plancher du camion et le sol du quai."},
 "À la réception, avec quel document": {"rep": "Le bon de livraison."},
 "Cite deux signes": {"rep": "Le marquage au sol hachuré, le panneau « Zone litiges — Ne pas stocker · Ne pas expédier », l'étiquette « BLOQUÉ » sur la palette (deux suffisent).",
   "note": "En couleur à l'écran : le marquage est rouge."},
 "Pourquoi la palette L2 est-elle bloquée": {"rep": "Il manque 2 cartons."},
 "Pourquoi une palette abîmée": {"pistes": [
   "On ne range pas un produit qu'on ne pourra pas vendre : il faudrait le chercher et le ressortir.",
   "Le fournisseur doit d'abord répondre (réserve, avoir, remplacement) : la palette attend à part.",
   "Si elle allait en stock, elle risquerait d'être envoyée à un client.",
   "Le stock à l'écran serait faux (des cartons comptés qui n'existent pas)."]},
 "T: N° | Mot | Ce que c’est": {"lignes": [[str(n), mot, df] for n, _, _, mot, df in MOTS],
   "note": "Mêmes mots et même photo qu'à l'écran (points 7 et 8 à la place du lot A des retours 5.3). Accepter toute définition juste."},
 "Pourquoi ne faut-il pas reconnaître une échelle": {"pistes": [
   "La couleur change d'un entrepôt à l'autre (bleu, gris, orange) : ce n'est pas elle qui fait l'échelle.",
   "On la reconnaît à sa forme : un montant vertical percé de trous, relié à un autre par des croisillons.",
   "C'est ce qui porte les lisses : sans elle, la lisse ne tient pas."]},
 "T: Lettre": None,
 "Combien de lisses compte cette travée": {"rep": "3 (en haut, au milieu, en bas ; celle du haut compte même vide)."},
 "Entre quoi et quoi va une travée": {"rep": "En largeur, d'une échelle à la suivante ; en hauteur, du sol jusqu'en haut du rack."},
 "Le bas de cette travée est un passage": {"rep": "Oui : elle va toujours d'une échelle à l'autre, du sol jusqu'en haut.",
   "note": "Décision 11 de Tristan : à dire en classe."},
 "Les barres qu’on voit au fond": {"rep": "Non : elles sont sur le rack de derrière ; elles ne sont pas accrochées aux échelles de devant."},
 "Bruno te dit « Va en travée 3 »": {"pistes": [
   "« Travée 3 » ne désigne qu'un seul endroit ; « la troisième étagère » dépend d'où on commence à compter.",
   "Tout le monde utilise les mêmes mots : pas d'erreur, pas de palette posée au mauvais endroit.",
   "On gagne du temps : pas besoin de chercher ni de demander."]},
 "T: Partie de l’adresse": {"lignes": [['A1', 'allée A, côté 1'], ['T03', 'travée 3'], ['N2', 'niveau 2'], ['E1', 'emplacement 1']],
   "note": "Jalon « adresse décomposée » : une seule validation à l'écran."},
 "Qu’as-tu trouvé à l’emplacement": {"rep": f"Une {NOMS[STOCK[ADRESSE][0]]} de {STOCK[ADRESSE][1]} kg.",
   "note": "Lu dans le stock de la plateforme (contenus/smoby-entrepot.js)."},
 "T: Adresse | Allée et côté": {"lignes": [[a] + decompose(a) for a in A_LIRE] + [[adresse(*x)] + decompose(adresse(*x)) for x in A_ECRIRE],
   "note": "Toutes ces adresses existent sur la plateforme de la séance : " + ' ; '.join(
       f"{a} = {NOMS[STOCK[a][0]]}, {STOCK[a][1]} kg" for a in A_LIRE + [adresse(*x) for x in A_ECRIRE]) + ". Attendre le zéro de la travée (T04, pas T4)."},
 "Pourquoi une adresse se lit-elle": {"pistes": [
   "C'est l'ordre dans lequel on marche : d'abord l'allée, puis la travée, puis on lève les yeux vers le niveau, puis on cherche l'emplacement.",
   "Comme une adresse postale à l'envers : pays, ville, rue, numéro… on réduit la zone à chaque étape.",
   "Si on commençait par l'emplacement (E1), il y en aurait des dizaines dans l'entrepôt."]},
 "Que fera Yanis cet après-midi": {"rep": "Il déchargera le premier camion de l'usine d'Arinthod, au quai n° 2."},
 "Par quoi Bruno dit-il": {"rep": "Par la sécurité."},
 "Tu as vu six endroits": {"pistes": [
   "L'aire de manœuvre et les quais : les camions reculent et voient mal derrière eux.",
   "L'allée principale : les chariots y roulent dans les deux sens.",
   "Les allées de stockage : une charge peut tomber d'en haut ; le cariste regarde ses fourches."]},
 "Qu'est-ce qui t'aidera le plus": {"pistes": [
   "Réponse personnelle. Au quai (ENT-5.4) : le chemin du quai à la zone de réception et la zone litiges ; au rangement (ENT-5.5) : les adresses et les niveaux."]},
}
del ENT_5_3["T: Lettre"]
for ph, mot in ESSENTIEL:
    ENT_5_3[ph.split('{}')[0].strip()] = {"rep": mot + '.'}
ENT_5_3['T: Mot | Définition (complète avec la banque de mots)'] = {
    "lignes": [[mot, m] for mot, _, m in LEXIQUE], "note": "Un mot de la banque par trou."}
corriges_data._DICOS['ENT-5.3'] = [ENT_5_3]

NOTIONS = [
    ["L'aire de manœuvre", "Circulation sur le site", "Piétons et engins sont séparés : un piéton ne circule pas où les camions manœuvrent ; il traverse au passage piétons."],
    ["Bruno le dit au quai", "Sécurité au quai", "Un camion est calé (et le niveleur en place) avant que le chariot entre dedans."],
    ["Un croisillon est tordu", "Sécurité des racks", "Un rack abîmé se signale tout de suite au responsable : il peut céder sous la charge."],
    ["Le sol est le niveau", "Adressage", "Le sol est le niveau 1 : N1."],
    ["Une travée, c'est", "Vocabulaire du rack", "Une travée est l'espace entre deux échelles, sur toute la hauteur du rack."],
    ["À l'adresse", "Adressage", "N1 = le sol ; N2 = le premier étage de lisses."],
    ["Où est écrite la charge maximale", "Sécurité des racks", "La plaque de charge, en bout de rack, donne la charge maximale par niveau (INRS ED 771)."],
    ["À pied dans l'allée principale", "Circulation sur le site", "À pied, on reste sur le côté et on traverse au passage piétons."],
]
T.finir('ENT-5.3', 'Smoby — la visite de la plateforme', 'ENT-5.3-smoby-visite-trame-eleve', NOTIONS,
        os.path.basename(__file__))
