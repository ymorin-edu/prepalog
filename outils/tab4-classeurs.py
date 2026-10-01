#!/usr/bin/env python3
"""TAB-4 — fabrique les treize classeurs de contenus/tab4/.

    python outils/tab4-classeurs.py

Les données viennent de outils/tab4-donnees.json, le MÊME fichier que lit le générateur de
contrôles (outils/tab4-exercices.mjs). Changer un chiffre là change le classeur et le
corrigé ensemble : ils ne peuvent pas se contredire. C'est la différence avec TAB-1, TAB-2
et TAB-3, où les classeurs venaient de la Suite Logistique et où modifier une donnée du
générateur ne touchait pas le fichier que l'élève télécharge.

Deux conséquences de les fabriquer plutôt que de les reprendre :

  - aucun corrigé caché ne peut voyager avec eux. Il n'y a pas d'onglet « Correction » à
    retirer, puisque rien ne vient d'ailleurs ;
  - les cellules de réponse sont vides par construction. Le test le revérifie quand même,
    parce qu'une construction juste aujourd'hui peut se casser demain.

Chaque classeur a deux onglets :

  « Consignes » — le geste Excel travaillé, l'objectif, le vocabulaire, le travail à faire ;
  « Exercice »  — un document professionnel : en-tête, tableau, indicateurs, pied de page.

AUCUNE FORMULE n'est écrite dans les classeurs : les cellules à remplir sont vides. C'est
l'élève qui écrit les formules, et c'est tout l'objet du module.
"""

import json
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RACINE = pathlib.Path(__file__).resolve().parent.parent
DONNEES = RACINE / 'outils' / 'tab4-donnees.json'
SORTIE = RACINE / 'contenus' / 'tab4'

# Charte : le vert d'Excel du site, l'encre sombre des classeurs de la Suite, et un beige
# pour les cases à compléter — même repère visuel que contenus/tab5-inventaire.xlsx.
ENCRE = '211F1A'
VERT = '107C41'
BEIGE = 'FBF3E0'
GRIS = '8A8577'
FILET = 'D8D4CC'
POLICE = 'Calibri'

f_titre = Font(name=POLICE, size=16, bold=True, color=VERT)
f_petit = Font(name=POLICE, size=9, color=GRIS)
f_doc = Font(name=POLICE, size=13, bold=True, color=ENCRE)
f_etiq = Font(name=POLICE, size=9, bold=True, color=GRIS)
f_val = Font(name=POLICE, size=10, color=ENCRE)
f_val_gras = Font(name=POLICE, size=10, bold=True, color=ENCRE)
f_entete = Font(name=POLICE, size=10, bold=True, color='FFFFFF')
f_corps = Font(name=POLICE, size=10, color=ENCRE)
f_note = Font(name=POLICE, size=9, italic=True, color=GRIS)
f_h = Font(name=POLICE, size=11, bold=True, color=ENCRE)
f_mono = Font(name='Consolas', size=11, bold=True, color=VERT)

fond_entete = PatternFill('solid', fgColor=ENCRE)
fond_remplir = PatternFill('solid', fgColor=BEIGE)

bord = Side(style='thin', color=FILET)
cadre = Border(left=bord, right=bord, top=bord, bottom=bord)
trait = Border(top=Side(style='thin', color=GRIS))

gauche = Alignment(horizontal='left', vertical='center')
centre = Alignment(horizontal='center', vertical='center')
droite = Alignment(horizontal='right', vertical='center')
hg = Alignment(horizontal='left', vertical='top', wrap_text=True)


def pose(ws, ref, valeur, police=None, fond=None, bordure=None, align=None, fmt=None):
    c = ws[ref]
    c.value = valeur
    if police:
        c.font = police
    if fond:
        c.fill = fond
    if bordure:
        c.border = bordure
    if align:
        c.alignment = align
    if fmt:
        c.number_format = fmt
    return c


def hauteur(texte, largeur_car):
    """Hauteur de ligne approchée pour un texte qui revient à la ligne."""
    return 14 + 13 * (len(texte) // largeur_car)


# --------------------------------------------------------------- onglet Consignes
def onglet_consignes(wb, ex, ent, numero, total):
    ws = wb.create_sheet('Consignes')
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 78
    ws.sheet_view.showGridLines = False

    pose(ws, 'B1', 'LYCÉE GASTON DARBOUX — NÎMES', f_petit)
    pose(ws, 'B2', 'TAB-4 — Une journée chez %s' % ent['nom'], f_titre)
    pose(ws, 'B3', 'CAP Opérateur logistique — Excel : les gestes de base', f_petit)
    pose(ws, 'B5', 'Exercice %d sur %d — %s' % (numero, total, ex['moment']), f_etiq)
    pose(ws, 'B6', ex['titre'], f_doc)

    l = 8
    pose(ws, 'B%d' % l, "Le geste Excel que vous apprenez", f_h)
    pose(ws, 'C%d' % l, ex['geste'], f_corps)
    ws['C%d' % l].alignment = hg
    ws.row_dimensions[l].height = hauteur(ex['geste'], 90)
    l += 1
    pose(ws, 'B%d' % l, 'À quoi ça sert dans le métier', f_h)
    pose(ws, 'C%d' % l, ex['objectif'], f_corps)
    ws['C%d' % l].alignment = hg
    ws.row_dimensions[l].height = hauteur(ex['objectif'], 90)

    l += 2
    pose(ws, 'B%d' % l, 'Vocabulaire à connaître', f_h)
    l += 1
    for mot, definition in ex['vocabulaire']:
        pose(ws, 'B%d' % l, mot, f_val_gras)
        ws['B%d' % l].alignment = hg
        pose(ws, 'C%d' % l, definition, f_corps)
        ws['C%d' % l].alignment = hg
        ws.row_dimensions[l].height = hauteur(definition, 88)
        l += 1

    l += 1
    pose(ws, 'B%d' % l, 'Ce que vous devez faire', f_h)
    l += 1
    for ligne in ex['consignes_rendues']:
        # Une ligne qui commence par des espaces est une formule ou une quantité à relever :
        # on la met en évidence, c'est ce que l'élève vient relire dix fois.
        formule = ligne.startswith('    ')
        pose(ws, 'B%d' % l, ligne.strip() if formule else ligne,
             f_mono if formule else f_corps)
        ws.merge_cells('B%d:C%d' % (l, l))
        ws['B%d' % l].alignment = hg
        ws.row_dimensions[l].height = hauteur(ligne, 125)
        l += 1

    l += 1
    pose(ws, 'B%d' % l, "Quand vous avez terminé : enregistrez le fichier, revenez sur Prepalog et "
                        "déposez-le. Vous pouvez le déposer autant de fois que vous voulez — "
                        "seul votre meilleur résultat est gardé.", f_note)
    ws.merge_cells('B%d:C%d' % (l, l))
    ws['B%d' % l].alignment = hg
    ws.row_dimensions[l].height = 28
    l += 2
    pose(ws, 'B%d' % l, ent['mention'], f_note)
    return ws


# ---------------------------------------------------------------- onglet Exercice
def entete_document(ws, ex, ent, blocs, nb):
    """L'en-tête d'un vrai document : qui écrit, à qui, quel document, quel numéro, quand."""
    doc = ex['document']
    fin = get_column_letter(nb)
    mid = get_column_letter(max(1, nb - 1))

    pose(ws, 'A1', ent['nom'], f_titre)
    pose(ws, 'A2', ent['activite'], f_petit)
    pose(ws, 'A3', '%s — %s' % (ent['adresse'], ent['ville']), f_petit)
    pose(ws, 'A4', 'Tél. %s — SIRET %s' % (ent['tel'], ent['siret']), f_petit)

    def a_droite(ligne, valeur, police):
        pose(ws, '%s%d' % (mid, ligne), valeur, police, align=droite)
        ws.merge_cells('%s%d:%s%d' % (mid, ligne, fin, ligne))

    a_droite(1, doc['type'], f_doc)
    a_droite(2, doc.get('numero', ''), f_val_gras)
    a_droite(3, doc.get('dateHeure', ''), f_petit)
    if doc.get('commande'):
        a_droite(4, 'Votre commande n° %s' % doc['commande'], f_petit)

    # Expéditeur et destinataire. Sur un bon de livraison, c'est la case qui dit de qui
    # vient la marchandise : un élève qui ne la regarde pas réclame au mauvais fournisseur.
    if doc.get('expediteur'):
        pose(ws, 'A6', 'ÉMIS PAR' if doc['expediteur'] == 'nous' else 'EXPÉDITEUR', f_etiq)
        for i, ligne in enumerate(blocs[doc['expediteur']]):
            pose(ws, 'A%d' % (7 + i), ligne, f_val_gras if i == 0 else f_val)
    if doc.get('destinataire'):
        etiq = 'LIVRÉ À' if doc['destinataire'] == 'nous' else 'DESTINATAIRE'
        pose(ws, '%s6' % mid, etiq, f_etiq, align=droite)
        ws.merge_cells('%s6:%s6' % (mid, fin))
        for i, ligne in enumerate(blocs[doc['destinataire']]):
            pose(ws, '%s%d' % (mid, 7 + i), ligne, f_val_gras if i == 0 else f_val, align=droite)
            ws.merge_cells('%s%d:%s%d' % (mid, 7 + i, fin, 7 + i))

    # La constante de conversion, quand il y en a une : une seule cellule, bien repérée,
    # que la formule de l'élève devra figer avec F4.
    if ex.get('constante'):
        lc = 9
        pose(ws, 'A%d' % lc, ex['constante']['libelle'], f_etiq)
        ws.merge_cells('A%d:B%d' % (lc, lc))
        pose(ws, 'C%d' % lc, ex['constante']['valeur'], f_val_gras, None, cadre, centre)

    for col in range(1, nb + 1):
        ws.cell(row=10, column=col).border = trait


def valeur_cellule(ligne, col, article):
    """La valeur d'une cellule de données : celle de la ligne, sinon celle du catalogue."""
    cle = col['cle']
    if cle in ligne:
        return ligne[cle]
    if article and cle in ('designation', 'unite', 'prix', 'poidsG'):
        return article[cle]
    return None


def bloc_indicateurs(ws, ex, ligne, nb, geo):
    """Les cases à calculer sous le tableau : libellé à gauche, case beige à droite.

    Les positions viennent de `geometrie` dans le JSON, que le générateur de contrôles lit
    aussi : le premier indicateur tombe `decalageIndicateurs` lignes sous la dernière ligne
    de données, dans la colonne `colonneIndicateurs`. Ne jamais calculer ça ici de son côté.
    """
    totaux = ex.get('totaux') or []
    if not totaux:
        return ligne
    colonne = geo['colonneIndicateurs']
    premiere = ligne + geo['decalageIndicateurs']
    pose(ws, 'A%d' % (premiere - 1), 'À CALCULER', f_etiq)
    # Le libellé occupe les colonnes A et B, la case est en C : la même géométrie pour les
    # treize classeurs, donc une seule chose à expliquer aux élèves.
    for i, t in enumerate(totaux):
        l = premiere + i
        pose(ws, 'A%d' % l, t['libelle'], f_val_gras)
        ws.merge_cells('A%d:B%d' % (l, l))
        ws['A%d' % l].alignment = gauche
        pose(ws, '%s%d' % (colonne, l), None, f_corps, fond_remplir, cadre, centre, t.get('format'))
    return premiere + len(totaux) - 1


def pied_document(ws, ex, ent, ligne, nb):
    fin = get_column_letter(nb)
    l = ligne + 2
    for col in range(1, nb + 1):
        ws.cell(row=l, column=col).border = trait
    l += 1
    if ex['document'].get('reserves'):
        pose(ws, 'A%d' % l, 'RÉSERVES À LA LIVRAISON', f_etiq)
        l += 1
        pose(ws, 'A%d' % l, "À remplir s'il manque quelque chose, ou si un colis est abîmé. "
                            "Rien d'écrit ici = marchandise acceptée telle quelle.", f_note)
        ws.merge_cells('A%d:%s%d' % (l, fin, l))
        l += 2
        pose(ws, 'A%d' % l, 'Nom du réceptionnaire :', f_val)
        pose(ws, '%s%d' % (get_column_letter(max(1, nb - 1)), l), 'Signature :', f_val)
        l += 2
    pose(ws, 'A%d' % l, ent['mention'], f_note)
    ws.merge_cells('A%d:%s%d' % (l, fin, l))


def onglet_exercice(wb, ex, ent, blocs, cat, geo):
    ws = wb.create_sheet('Exercice')
    tab = ex['tableau']
    cols = tab['colonnes']
    nb = len(cols)
    ws.sheet_view.showGridLines = False

    for i, col in enumerate(cols):
        ws.column_dimensions[get_column_letter(i + 1)].width = col.get('largeur', 14)

    entete_document(ws, ex, ent, blocs, nb)

    lig = geo['ligneEntete']
    for i, col in enumerate(cols):
        pose(ws, '%s%d' % (get_column_letter(i + 1), lig), col['titre'],
             f_entete, fond_entete, cadre, centre)
    ws.row_dimensions[lig].height = 22

    for j, ligne in enumerate(tab['lignes']):
        r = geo['premiereLigne'] + j
        article = cat.get(ligne.get('ref'))
        for i, col in enumerate(cols):
            ref = '%s%d' % (get_column_letter(i + 1), r)
            if col.get('aRemplir'):
                # Case à compléter : vide, sur fond beige. Aucune formule n'est posée.
                pose(ws, ref, None, f_corps, fond_remplir, cadre, centre, col.get('format'))
                continue
            v = valeur_cellule(ligne, col, article)
            pose(ws, ref, v, f_corps, None, cadre,
                 gauche if isinstance(v, str) else centre, col.get('format'))

    derniere = geo['premiereLigne'] + len(tab['lignes']) - 1
    derniere = bloc_indicateurs(ws, ex, derniere, nb, geo)
    pied_document(ws, ex, ent, derniere, nb)
    return ws


def rendre_consignes(ex, cat):
    """Développe LISTE_COMPTAGE en la liste réelle des quantités comptées."""
    sorties = []
    for ligne in ex['consignes']:
        if ligne != 'LISTE_COMPTAGE':
            sorties.append(ligne)
            continue
        col = ex['calcul']['colonne']
        for l in ex['tableau']['lignes']:
            art = cat[l['ref']]
            unite = art['unitePluriel'] if l[col] > 1 else art['unite']
            # La désignation garde sa casse : « S3 T42 » n'est pas du texte courant.
            sorties.append('    %s %s — %s  (%s)' % (l[col], unite, art['designation'], l['ref']))
    return sorties


def main():
    d = json.loads(DONNEES.read_text(encoding='utf-8'))
    ent, geo = d['entreprise'], d['geometrie']
    cat = {a['ref']: a for a in d['catalogue']}
    blocs = {'nous': d['nous'], 'fournisseur': d['fournisseur'], 'transporteur': d['transporteur']}
    SORTIE.mkdir(parents=True, exist_ok=True)
    total = len(d['exercices'])

    for numero, ex in enumerate(d['exercices'], start=1):
        ex['consignes_rendues'] = rendre_consignes(ex, cat)
        wb = Workbook()
        wb.remove(wb.active)
        onglet_consignes(wb, ex, ent, numero, total)
        onglet_exercice(wb, ex, ent, blocs, cat, geo)
        wb.properties.title = '%s — %s' % (ex['id'], ex['titre'])
        wb.properties.creator = 'Prepalog'
        chemin = SORTIE / ('%s-%s.xlsx' % (ex['id'], ex['fichier']))
        wb.save(chemin)
        print('  %s' % chemin.relative_to(RACINE))

    print('%d classeurs écrits dans contenus/tab4/.' % total)


if __name__ == '__main__':
    main()
